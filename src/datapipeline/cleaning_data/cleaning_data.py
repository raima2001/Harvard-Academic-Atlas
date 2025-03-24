from fpdf import FPDF
from docx import Document
import pandas as pd


data_df = pd.read_csv('/app/data/courses_data_all_pages.csv')

data_df.head()

data_df.info()

data_df.describe(include='all')

data_df.isnull().sum()

data_df['Prerequisite'].fillna('prerequisites not provided', inplace=True)
data_df['Time'].fillna(0, inplace=True)
data_df['Exam Type'].fillna('exam type not given', inplace=True)
data_df['Description'].fillna('description unavailable', inplace=True)
data_df['Extra Note'].fillna('no further info', inplace=True)
data_df['Room'].fillna('room number not provided', inplace=True)
data_df['Areas of Interest'].fillna('No info given', inplace=True)
data_df['Instructors'].fillna('No info given', inplace=True)

data_df['Time']

import re

day_mapping = {
    'M': 'Monday',
    'T': 'Tuesday',
    'W': 'Wednesday',
    'Th': 'Thursday',
    'F': 'Friday',
    'Sa': 'Saturday',
    'Su': 'Sunday'
}

def expand_days(days):
    if pd.isna(days):
        return 'days not provided'
    days_list = re.split(r',\s*', days.strip())  # Split on comma with optional spaces
    full_days = [day_mapping.get(day.strip(), day.strip()) for day in days_list]
    return ', '.join(full_days)

data_df['Days'] = data_df['Time'].str.extract(r'^([A-Za-z, ]+)', expand=False)
data_df['Hours'] = data_df['Time'].str.extract(r'(\d{1,2}:\d{2}[ap]m – \d{1,2}:\d{2}[ap]m)', expand=False)

data_df['Days'] = data_df['Days'].apply(expand_days)

data_df.drop('Time', axis=1, inplace=True)

data_df['Credits']

data_df.duplicated().sum()

data_df.dtypes

columns_to_convert = data_df.columns.difference(['Hours'])

data_df[columns_to_convert] = data_df[columns_to_convert].astype('string')

data_df.dtypes

data_df[['Offered', 'Course Type']] = data_df['Type'].str.extract(r'(\w+ \d{4}) (.*)', expand=True)

data_df.drop('Type', axis=1, inplace=True)

data_df['Credits'] = data_df['Credits'].str.extract(r'(\d+)').astype(int)

data_df.dtypes

data_df.info()
data_df.describe(include='all')

data_df.to_csv(r'D:\ac215_masalachai\data\HLS_cleaned_data.csv', index=False)

def export_to_pdf(df, pdf_path):
    pdf = FPDF()
    pdf.set_font("Arial", size=10)
    pdf.add_page()

    # Add header row
    for col in df.columns:
        pdf.cell(40, 10, col, border=1)
    pdf.ln()

    # Add data rows
    for _, row in df.iterrows():
        for item in row:
            pdf.cell(40, 10, str(item), border=1)
        pdf.ln()

    pdf.output(pdf_path)

def export_to_docx(df, docx_path):
    doc = Document()
    doc.add_heading('Courses Data', 0)

    # Add table
    table = doc.add_table(rows=1, cols=len(df.columns))
    hdr_cells = table.rows[0].cells
    for idx, col in enumerate(df.columns):
        hdr_cells[idx].text = col

    for _, row in df.iterrows():
        row_cells = table.add_row().cells
        for idx, item in enumerate(row):
            row_cells[idx].text = str(item)

    doc.save(docx_path)

output_pdf_path = r'D:\ac215_masalachai\data\HLS_cleaned_data.pdf'
output_docx_path = r'D:\ac215_masalachai\data\HLS_cleaned_data.docx'
export_to_pdf(data_df, output_pdf_path)
export_to_docx(data_df, output_docx_path)

print("Data exported successfully to CSV, PDF, and DOCX formats.")
